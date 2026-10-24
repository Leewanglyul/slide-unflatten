import os
import sys

import pytest
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def draw_slide(w=1280, h=720):
    """A synthetic flattened slide: title, two bullet lines, a blue blob 'illustration'. Returns (img, lines)."""
    from unflatten.types import Line
    img = Image.new("RGB", (w, h), (250, 247, 240))
    d = ImageDraw.Draw(img)
    d.ellipse((900, 300, 1180, 600), fill=(60, 120, 220))            # object
    d.rectangle((880, 120, 1200, 200), fill=(250, 200, 80))           # second object
    lines = []
    specs = [("Polar Frontiers Report", 80, 70, 56, FONT_B, (30, 30, 60)),
             ("• Ice is melting faster than expected", 100, 260, 30, FONT, (40, 40, 40)),
             ("• Shipping routes are opening", 100, 320, 30, FONT, (40, 40, 40))]
    for text, x, y, size, fpath, col in specs:
        f = ImageFont.truetype(fpath, size)
        d.text((x, y), text, fill=col, font=f)
        l, t, r, b = d.textbbox((x, y), text, font=f)
        lines.append(Line(box=(int(l), int(t), int(r) + 1, int(b) + 1), text=text, conf=0.95))
    return img, lines


class FakeOCR:
    name = "fake"

    def __init__(self, lines):
        self._lines = lines
        self.calls = 0

    def lines(self, img):
        self.calls += 1
        return list(self._lines)


@pytest.fixture
def flat_png(tmp_path):
    img, lines = draw_slide()
    p = tmp_path / "flat.png"
    img.save(p)
    return str(p), lines


@pytest.fixture
def pptx_mixed(tmp_path):
    """slide 1 flattened (full picture + notes), slide 2 real text + small pic, slide 3 title only."""
    from pptx import Presentation
    from pptx.util import Inches
    img, lines = draw_slide()
    ip = tmp_path / "s1.png"
    img.save(ip)
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    s1.shapes.add_picture(str(ip), 0, 0, prs.slide_width, prs.slide_height)
    s1.notes_slide.notes_text_frame.text = "speaker notes 1"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Real editable title"
    s2.shapes.add_picture(str(ip), Inches(1), Inches(2), Inches(4), Inches(2))
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "only text"
    p = tmp_path / "mixed.pptx"
    prs.save(p)
    return str(p), lines


@pytest.fixture
def pdf_text(tmp_path):
    """2-page PDF: page 1 has a text layer + embedded image, page 2 is a pure raster (no text)."""
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=960, height=540)
    page.insert_text((50, 80), "Deck Title From PDF", fontsize=36, fontname="hebo", color=(0.1, 0.2, 0.3))
    page.insert_text((50, 160), "First body line in the pdf", fontsize=18, fontname="helv")
    page.insert_text((50, 190), "Second body line in the pdf", fontsize=18, fontname="helv")
    img, _ = draw_slide(400, 300)
    ip = tmp_path / "emb.png"
    img.save(ip)
    page.insert_image(pymupdf.Rect(600, 250, 900, 475), filename=str(ip))
    page2 = doc.new_page(width=960, height=540)
    big, _ = draw_slide()
    bp = tmp_path / "raster.png"
    big.save(bp)
    page2.insert_image(page2.rect, filename=str(bp))
    p = tmp_path / "text.pdf"
    doc.save(p)
    return str(p)
