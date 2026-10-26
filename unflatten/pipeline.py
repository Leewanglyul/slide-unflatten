"""Orchestration: load → recognise → erase → split → assemble.

Slides stream through one at a time (bounded look-ahead), so memory stays flat for 100-page PDFs:
OCR runs on the main thread, each slide's Claude call is submitted as soon as its OCR finishes and runs on
a thread while the next slides are OCR'd / the previous ones are erased and split. Results go straight to
the per-slide cache; assembly reads them back one slide at a time."""
from __future__ import annotations

import json
import os
import time
import traceback
import zipfile
from collections import deque
from concurrent.futures import Future, ThreadPoolExecutor
from dataclasses import dataclass, field
from typing import Callable, Dict, List, Optional

import numpy as np

from . import builder
from .cache import SlideCache
from .config import Config, estimate_cost_usd, parse_slide_range, USD_KRW
from .fonts import bundle_fonts, plan_fonts
from .i18n import t
from .imageops import classify_background, dominant_colors
from .inputs import Deck, SlideInput, load_deck
from .inpaint import box_mask, inpaint, text_mask
from .layout import FontMeasurer, blocks_from_claude, finalize_blocks, group_heuristic, refit_fonts
from .ocr import NullOCRBackend, easyocr_available, make_backend
from .report import render_preview, side_by_side, write_html_report
from .segment import pieces_from_embedded, split_semantic, split_simple
from .types import Piece, SlideAnalysis, SlideSummary

Progress = Callable[[int, int, str], None]
SEC_PER_SLIDE_GPU = 15
SEC_PER_SLIDE_CPU = 90


@dataclass
class Result:
    out_path: str
    manifest_path: str
    report_path: Optional[str]
    summaries: List[SlideSummary]
    previews: List[Optional[str]] = field(default_factory=list)     # image paths, aligned with summaries
    cost_usd: float = 0.0
    input_tokens: int = 0
    output_tokens: int = 0
    seconds: float = 0.0
    used_claude: bool = False
    font_files: List[str] = field(default_factory=list)

    def cost_line(self, lang="ko") -> str:
        if self.input_tokens == 0 and not self.used_claude:
            return t(lang, "warn_noclaude")
        return t(lang, "cost", inp=self.input_tokens, out=self.output_tokens, usd=self.cost_usd)


# ───────────── estimate (metadata only: no page is rendered) ─────────────
def estimate(path: str, cfg: Optional[Config] = None) -> Dict:
    cfg = cfg or Config()
    deck = load_deck(path, dpi=cfg.pdf_dpi)
    wanted = set(parse_slide_range(cfg.slide_range, len(deck.slides)))
    flat = [s for s in deck.slides if s.kind == "flat" and s.index in wanted]
    text_layer = [s for s in deck.slides if s.kind == "text_layer" and s.index in wanted]
    passthrough = [s for s in deck.slides if s.kind == "passthrough"]
    gpu = False
    try:
        import torch
        gpu = torch.cuda.is_available()
    except Exception:
        pass
    per = SEC_PER_SLIDE_GPU if gpu else SEC_PER_SLIDE_CPU
    if cfg.mode == "text":
        per *= 0.5
    minutes = max(1, round((len(flat) * per + len(text_layer) * 2) / 60))
    usd = estimate_cost_usd(len(flat), cfg.model) if (cfg.use_claude and cfg.resolved_api_key()) else 0.0
    return {"n_slides": len(deck.slides), "n_flat": len(flat), "n_text_layer": len(text_layer),
            "n_pass": len(passthrough), "minutes": minutes, "usd": usd, "krw": usd * USD_KRW, "gpu": gpu,
            "line": t(cfg.ui_lang, "estimate", n=len(flat) + len(text_layer), min=minutes, usd=usd, krw=usd * USD_KRW)}


# ───────────── per-slide steps ─────────────
def _recognise(si: SlideInput, ocr) -> SlideAnalysis:
    if si.kind == "text_layer":
        lines = si.lines or []
        return SlideAnalysis(lines=lines, blocks=group_heuristic(lines, img_h=si.image.height), objects=[],
                             background_type=classify_background(si.image), source="text_layer")
    return SlideAnalysis(lines=ocr.lines(si.image), blocks=[], objects=[], background_type="unknown", source=ocr.name)


def _analyse_with_claude(si: SlideInput, analysis: SlideAnalysis, cfg: Config, service):
    out = service.analyze(si.image, analysis.lines, cfg.object_detail)
    return None if out is None else blocks_from_claude(out, analysis.lines)


def _erase_and_split(si: SlideInput, analysis: SlideAnalysis, cfg: Config, log):
    img = si.image
    arr = np.array(img)
    kept = sorted({i for b in analysis.blocks for i in b.lines})
    clean = si.clean_image if si.kind == "text_layer" else None
    if clean is None:
        log("  " + t(cfg.ui_lang, "erase", n=len(kept)))
        clean = inpaint(img, text_mask(arr.shape, [analysis.lines[i] for i in kept], arr, cfg.mask_pad), log)
    if cfg.remove_watermark and analysis.watermarks:
        log("  " + t(cfg.ui_lang, "watermark", n=len(analysis.watermarks)))
        clean = inpaint(clean, box_mask(arr.shape, analysis.watermarks), log)
    pieces: List[Piece] = []
    background = clean
    if si.kind == "text_layer" and si.embedded_images and cfg.split_objects:
        pieces = pieces_from_embedded(si.embedded_images)
    elif cfg.split_objects:
        if analysis.background_type == "photo":
            log("  " + t(cfg.ui_lang, "split_skip_photo"))
        else:
            if cfg.object_detail in ("coarse", "fine") and analysis.objects:
                log("  " + t(cfg.ui_lang, "split_sam", names=", ".join(o["name"] for o in analysis.objects)))
                try:
                    background, pieces = split_semantic(clean, analysis.objects, cfg, log)
                except Exception as e:
                    log("  " + t(cfg.ui_lang, "sam_fail", err=str(e)[:120]))
            if not pieces:
                log("  " + t(cfg.ui_lang, "split_simple"))
                background, pieces = split_simple(clean, cfg, log)
    return background, pieces


def _resolve_backend_name(cfg: Config, service, ocr_override) -> str:
    if ocr_override is not None:
        return ocr_override.name
    if cfg.ocr_backend in ("auto", "easyocr") and easyocr_available():
        return "easyocr"
    if cfg.ocr_backend in ("auto", "claude") and service is not None:
        return "claude"
    return "none"


# ───────────── main ─────────────
def convert(path: str, cfg: Optional[Config] = None, out_path: Optional[str] = None,
            progress: Optional[Progress] = None, log=print, ocr_override=None,
            known_keys: Optional[Dict[int, str]] = None) -> Result:
    """`known_keys`: slide index -> cache key from a previous run (manifest); those slides are reused as-is
    even if the backends available now differ (that is what `rebuild` wants)."""
    cfg = cfg or Config()
    known_keys = known_keys or {}
    t0 = time.time()
    lang = cfg.ui_lang
    progress = progress or (lambda d, n, m: None)

    deck = load_deck(path, dpi=cfg.pdf_dpi)
    wanted = set(parse_slide_range(cfg.slide_range, len(deck.slides)))
    out_dir = cfg.out_dir or os.path.dirname(os.path.abspath(path))
    os.makedirs(out_dir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(path))[0]
    out_path = out_path or os.path.join(out_dir, f"{stem}_editable.pptx")
    cache = SlideCache(cfg.cache_dir or os.path.join(out_dir, ".unflatten-cache"))
    preview_dir = os.path.join(out_dir, f"{stem}_previews")

    # services (resolved up front so the cache key reflects what really runs)
    service = None
    if cfg.use_claude and cfg.resolved_api_key():
        try:
            from .claude_client import ClaudeService
            service = ClaudeService(cfg)
            ok, msg = service.validate()
            if not ok:
                log(t(lang, "claude_fail", err=msg))
                service = None
        except Exception as e:
            log(t(lang, "claude_fail", err=str(e)[:160]))
            service = None
    backend_name = _resolve_backend_name(cfg, service, ocr_override)
    ocr = ocr_override                                   # instantiated lazily (EasyOCR loads models)
    plan = plan_fonts(cfg)
    measurer = FontMeasurer(cfg)
    pool = ThreadPoolExecutor(max_workers=max(1, cfg.workers)) if service is not None else None
    state = {"service": service, "n_done": 0}

    work = [s for s in deck.slides if s.kind in ("flat", "text_layer") and s.index in wanted]
    total = len(work)
    outcomes: Dict[int, dict] = {}                       # index -> {"key", "ok"}
    summaries: Dict[int, SlideSummary] = {}
    previews: Dict[int, Optional[str]] = {}
    seen: Dict[str, int] = {}                            # cache key -> first slide index in this run
    inflight: deque = deque()

    def finish(item: dict):
        si: SlideInput = item["si"]
        key = item["key"]
        summ = SlideSummary(index=si.index, status="ok", cache_key=key)
        ts = time.time()
        try:
            if item.get("cached") or item.get("dup"):
                if not cache.has(key):
                    raise RuntimeError("duplicate of a slide that failed")
                if cache.split_params(key) != cfg.split_params():
                    # only edge/cut settings changed: reuse OCR + Claude, redo erase and split
                    analysis = cache.load_analysis(key)
                    log("  " + t(lang, "resplit"))
                    background, pieces = _erase_and_split(si, analysis, cfg, log)
                    cache.save(key, analysis, background, pieces, cfg.split_params())
                else:
                    analysis, background, pieces = cache.load(key)
                save_key = key
            else:
                analysis: SlideAnalysis = item["analysis"]
                fut: Optional[Future] = item.get("future")
                if fut is not None:
                    progress(state["n_done"], total, f"[{si.index}] {t(lang, 'claude')}")
                    try:
                        got = fut.result()
                    except Exception as e:
                        got = None
                        log("  " + t(lang, "claude_fail", err=f"{type(e).__name__}: {str(e)[:120]}"))
                        if type(e).__name__ in ("AuthenticationError", "PermissionDeniedError"):
                            log("  " + t(lang, "claude_auth"))
                            state["service"] = None
                            for other in inflight:       # don't keep paying for calls that will fail
                                f2 = other.get("future")
                                if f2 is not None:
                                    f2.cancel()
                    if got is not None:
                        analysis.blocks, analysis.objects, analysis.background_type, analysis.watermarks = got
                        analysis.used_claude = True
                if not analysis.blocks:
                    analysis.blocks = group_heuristic(analysis.lines, img_h=si.image.height)
                if analysis.background_type == "unknown":
                    analysis.background_type = classify_background(si.image)
                analysis.blocks = finalize_blocks(analysis.blocks, analysis.lines, np.array(si.image), measurer)
                progress(state["n_done"], total, f"[{si.index}] {t(lang, 'erase', n=len(analysis.lines))}")
                background, pieces = _erase_and_split(si, analysis, cfg, log)
                # a slide that fell back to heuristics is cached under the heuristic key, never the Claude key
                degraded = si.kind == "flat" and service is not None and not analysis.used_claude
                save_key = SlideCache.key(si.image, cfg, backend_name, False) if degraded else key
                cache.save(save_key, analysis, background, pieces, cfg.split_params())
            summ.cache_key = save_key
            summ.status = "text_layer" if si.kind == "text_layer" else "ok"
            summ.n_lines, summ.n_blocks, summ.n_pieces = len(analysis.lines), len(analysis.blocks), len(pieces)
            summ.low_conf = sum(1 for l in analysis.lines if l.conf < 0.5)
            summ.texts = [b.text.replace("\n", " ⏎ ") for b in analysis.blocks]
            if summ.low_conf:
                summ.warnings.append(t(lang, "warn_lowconf", n=summ.low_conf))
            if si.kind == "flat" and not analysis.used_claude:
                summ.warnings.append(t(lang, "warn_noclaude"))
            if si.kind == "flat" and not analysis.lines:
                summ.warnings.append("no text recognised")
            outcomes[si.index] = {"key": save_key, "ok": True}
            if cfg.make_report:
                os.makedirs(preview_dir, exist_ok=True)
                pv_path = os.path.join(preview_dir, f"{si.index:03d}.jpg")
                side_by_side(si.image, render_preview(background, pieces, analysis.blocks, measurer)).save(pv_path, quality=80)
                previews[si.index] = pv_path
        except Exception as e:
            log(t(lang, "slide_fail", i=si.index, err=f"{type(e).__name__}: {e}"))
            log(traceback.format_exc(limit=3))
            summ.status = "failed"
            summ.warnings.append(f"{type(e).__name__}: {str(e)[:100]}")
            outcomes[si.index] = {"key": key, "ok": False}
            if cfg.make_report:
                try:
                    os.makedirs(preview_dir, exist_ok=True)
                    pv_path = os.path.join(preview_dir, f"{si.index:03d}.jpg")
                    side_by_side(si.image, si.image).save(pv_path, quality=80)
                    previews[si.index] = pv_path
                except Exception:
                    previews[si.index] = None
        finally:
            si.release()
        summ.seconds = time.time() - ts
        summaries[si.index] = summ
        state["n_done"] += 1
        progress(state["n_done"], total, f"[{si.index}] {t(lang, 'done', path='').rstrip(': ')}")

    # ─── streaming recognition with bounded look-ahead ───
    for si in work:
        key = known_keys.get(si.index)
        if not key or not cache.has(key):
            key = SlideCache.key(si.image, cfg, backend_name, state["service"] is not None, cfg.model)
        item = {"si": si, "key": key}
        if cache.has(key):
            item["cached"] = True
            log(f"[{si.index}/{len(deck.slides)}] " + t(lang, "cached"))
        elif key in seen:
            item["dup"] = True
        else:
            stage = "text_layer" if si.kind == "text_layer" else "ocr"
            progress(state["n_done"], total, f"[{si.index}] {t(lang, stage)}")
            log(f"[{si.index}/{len(deck.slides)}] " + t(lang, stage))
            if ocr is None and si.kind == "flat":
                ocr = make_backend(cfg, state["service"], log)
            try:
                item["analysis"] = _recognise(si, ocr or NullOCRBackend())
                if si.kind == "flat" and state["service"] is not None and pool is not None:
                    item["future"] = pool.submit(_analyse_with_claude, si, item["analysis"], cfg, state["service"])
            except Exception as e:
                item["error"] = e
                item["analysis"] = None
        seen.setdefault(key, si.index)
        inflight.append(item)
        while len(inflight) > max(1, cfg.workers):
            nxt = inflight.popleft()
            if nxt.get("analysis", 0) is None and "error" in nxt:
                _fail(nxt, outcomes, summaries, previews, lang, log)
            else:
                finish(nxt)
    while inflight:
        nxt = inflight.popleft()
        if nxt.get("analysis", 0) is None and "error" in nxt:
            _fail(nxt, outcomes, summaries, previews, lang, log)
        else:
            finish(nxt)
    if pool:
        pool.shutdown(wait=False)

    # ─── assembly ───
    _assemble(deck, outcomes, wanted, cfg, plan, measurer, cache, out_path)

    for s in deck.slides:
        if s.kind == "passthrough":
            summaries[s.index] = SlideSummary(index=s.index, status="passthrough", warnings=[s.reason] if s.reason else [])
        elif s.index not in wanted:
            summaries[s.index] = SlideSummary(index=s.index, status="skipped")
    ordered = [summaries[i] for i in sorted(summaries)]
    pv_list = [previews.get(i) for i in sorted(summaries)]

    res = Result(out_path=out_path, manifest_path=os.path.join(out_dir, f"{stem}_manifest.json"), report_path=None,
                 summaries=ordered, previews=pv_list, seconds=time.time() - t0, used_claude=service is not None)
    if service is not None:
        res.cost_usd, res.input_tokens, res.output_tokens = service.cost_usd, service.input_tokens, service.output_tokens
    if cfg.bundle_fonts:
        res.font_files = bundle_fonts(plan, out_dir)
    if cfg.make_report:
        res.report_path = write_html_report(os.path.join(out_dir, f"{stem}_report.html"), f"{stem} — unflatten report",
                                            ordered, pv_list, res.cost_line(lang), lang)
        log(t(lang, "report", path=res.report_path))
    with open(res.manifest_path, "w", encoding="utf-8") as f:
        json.dump({"source": os.path.abspath(path), "output": os.path.abspath(out_path), "config": cfg.to_dict(),
                   "cache_dir": cache.root, "slides": [{"index": s.index, "status": s.status, "cache_key": s.cache_key,
                                                        "texts": s.texts} for s in ordered]},
                  f, ensure_ascii=False, indent=1)
    log(t(lang, "done", path=out_path))
    if res.input_tokens:
        log(res.cost_line(lang))
    return res


def _fail(item, outcomes, summaries, previews, lang, log):
    si = item["si"]
    e = item["error"]
    log(t(lang, "slide_fail", i=si.index, err=f"{type(e).__name__}: {e}"))
    outcomes[si.index] = {"key": item["key"], "ok": False}
    summaries[si.index] = SlideSummary(index=si.index, status="failed", warnings=[f"{type(e).__name__}: {str(e)[:100]}"])
    previews[si.index] = None
    si.release()


def _assemble(deck: Deck, outcomes, wanted, cfg: Config, plan, measurer, cache: SlideCache, out_path: str):
    from pptx import Presentation
    sw, sh = deck.width_emu, deck.height_emu
    has_pass = any(s.kind == "passthrough" for s in deck.slides)
    if deck.kind == "pptx":
        prs = deck.source_prs
        theme_ok = not has_pass          # don't re-font a deck that still has the author's own slides
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = sw, sh
        theme_ok = True
    if cfg.use_theme_fonts and theme_ok:
        sample = [s for s in deck.slides if s.kind != "passthrough"][:12]
        colors = dominant_colors([s.image for s in sample]) if sample else []
        for s in sample:
            s.release()
        builder.apply_theme(prs, plan, colors)
    eff_cfg = cfg if (cfg.use_theme_fonts and theme_ok) else _with(cfg, use_theme_fonts=False)

    for si in deck.slides:
        if deck.kind == "pptx":
            slide = prs.slides[si.index - 1]
            if si.kind == "passthrough" or si.index not in wanted:
                continue
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            if si.notes:
                slide.notes_slide.notes_text_frame.text = si.notes
            if si.index not in wanted:
                builder.add_failed_slide(slide, si.image, si.rect_emu, (sw, sh), cfg)
                si.release()
                continue
        oc = outcomes.get(si.index) or {"ok": False}
        if not oc["ok"]:
            if deck.kind != "pptx":          # pptx: the original picture shape is simply left in place
                builder.add_failed_slide(slide, si.image, si.rect_emu, (sw, sh), cfg)
                si.release()
            continue
        analysis, background, pieces = cache.load(oc["key"])
        refit_fonts(analysis.blocks, measurer)
        if deck.kind == "pptx":
            builder.remove_empty_placeholders(slide)
            builder.build_slide(slide, None, background, pieces, analysis.blocks, eff_cfg, plan, si.rect_emu, (sw, sh),
                                pdf_font_names=False, orig_shape_id=si.pic_shape_id)
        else:
            builder.build_slide(slide, si.image if cfg.keep_hidden_original else None, background, pieces,
                                analysis.blocks, eff_cfg, plan, si.rect_emu, (sw, sh), pdf_font_names=(si.kind == "text_layer"))
            si.release()
    prs.save(out_path)


def _with(cfg: Config, **kw) -> Config:
    d = cfg.to_dict()
    d.update(kw)
    c = Config.from_dict(d)
    c.api_key = cfg.api_key
    return c


# ───────────── rebuild from manifest (assembly-only changes: fonts, grouping, wrap…) ─────────────
def rebuild(manifest_path: str, overrides: Optional[dict] = None, progress: Optional[Progress] = None,
            log=print) -> Result:
    with open(manifest_path, encoding="utf-8") as f:
        man = json.load(f)
    cfg = Config.from_dict(man["config"])
    cfg.cache_dir = man.get("cache_dir") or cfg.cache_dir
    for k, v in (overrides or {}).items():
        if k in Config.__dataclass_fields__:
            setattr(cfg, k, v)
    keys = {int(s["index"]): s["cache_key"] for s in man.get("slides", []) if s.get("cache_key")}
    return convert(man["source"], cfg, out_path=man.get("output"), progress=progress, log=log, known_keys=keys)


def bundle_zip(res: Result) -> str:
    zpath = os.path.splitext(res.out_path)[0] + "_bundle.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(res.out_path, os.path.basename(res.out_path))
        if res.report_path and os.path.exists(res.report_path):
            z.write(res.report_path, os.path.basename(res.report_path))
        for fpath in res.font_files:
            z.write(fpath, "fonts/" + os.path.basename(fpath))
        if os.path.exists(res.manifest_path):
            z.write(res.manifest_path, os.path.basename(res.manifest_path))
    return zpath
