"""Assemble an editable slide: hidden original, background, picture pieces, text boxes.

Shapes are grouped ("pictures" / "text"), named and given alt text; the title becomes a real title
placeholder; runs reference theme fonts so one change in the slide master re-fonts the deck."""
from __future__ import annotations

import re
from typing import List, Optional, Sequence, Tuple

from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .fonts import FontPlan
from .imageops import to_jpg, to_png, upscale
from .types import Block, Piece

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
LABELS = {
    "ko": {"orig": "원본(숨김)", "bg": "배경", "pics": "그림", "texts": "텍스트", "pic": "그림", "text": "텍스트"},
    "en": {"orig": "original (hidden)", "bg": "background", "pics": "pictures", "texts": "text", "pic": "picture", "text": "text"},
}


def _L(cfg, key):
    return LABELS.get(cfg.ui_lang, LABELS["en"])[key]


# ───────────── theme ─────────────
def apply_theme(prs, plan: FontPlan, accent_colors: Sequence[Tuple[int, int, int]] = ()):
    """Set theme major/minor fonts (title/body) and accent colours on the first slide master."""
    try:
        theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    except Exception:
        return False
    root = etree.fromstring(theme_part.blob)
    ns = {"a": NS_A}
    for tag, name in (("majorFont", plan.bold_name), ("minorFont", plan.regular_name)):
        for el in root.findall(f".//a:fontScheme/a:{tag}", ns):
            for sub in ("latin", "ea", "cs"):
                node = el.find(f"a:{sub}", ns)
                if node is None:
                    node = etree.SubElement(el, f"{{{NS_A}}}{sub}")
                node.set("typeface", name)
    for i, rgb in enumerate(list(accent_colors)[:6], 1):
        acc = root.find(f".//a:clrScheme/a:accent{i}", ns)
        if acc is None:
            continue
        for child in list(acc):
            acc.remove(child)
        c = etree.SubElement(acc, f"{{{NS_A}}}srgbClr")
        c.set("val", "%02X%02X%02X" % tuple(rgb))
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return True


# ───────────── helpers ─────────────
def iter_shapes(shapes):
    """Walk shapes recursively (group shapes included)."""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # GROUP
            yield from iter_shapes(sh.shapes)


def _alt(shape, text: str):
    shape._element._nvXxPr.cNvPr.set("descr", text)


def _hide(shape):
    shape._element._nvXxPr.cNvPr.set("hidden", "1")


def _make_title_placeholder(shape):
    nvPr = shape._element.nvSpPr.nvPr
    if nvPr.find(qn("p:ph")) is None:
        ph = etree.SubElement(nvPr, qn("p:ph"))
        ph.set("type", "title")


def remove_empty_placeholders(slide):
    for sh in list(slide.shapes):
        if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip():
            sh._element.getparent().remove(sh._element)


def remove_shape_by_id(slide, shape_id: int):
    for sh in list(slide.shapes):
        if sh.shape_id == shape_id:
            sh._element.getparent().remove(sh._element)
            return True
    return False


def _add_bullet(p, font_pt: float):
    pPr = p._p.get_or_add_pPr()
    mar = int(font_pt * 1.1 * 12700)
    pPr.set("marL", str(mar))
    pPr.set("indent", str(-mar))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def _segments(text: str, highlights: List[str]) -> List[Tuple[str, bool]]:
    if not highlights:
        return [(text, False)]
    words = sorted({h.strip() for h in highlights if h.strip()}, key=len, reverse=True)
    if not words:
        return [(text, False)]
    pat = re.compile("|".join(re.escape(w) for w in words), re.IGNORECASE)
    out, pos = [], 0
    for m in pat.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((m.group(0), True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out or [(text, False)]


def _style_run(run, size_pt: float, bold: bool, color, font_name: str, theme_ref: Optional[str], lang: str,
               highlight_hex: Optional[str] = None):
    f = run.font
    f.size = Pt(round(size_pt, 1))
    f.bold = bold
    f.color.rgb = RGBColor(*color)
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", lang)
    if highlight_hex:
        hl = etree.SubElement(rPr, qn("a:highlight"))   # sequence: after fill/effects, before latin
        c = etree.SubElement(hl, qn("a:srgbClr"))
        c.set("val", highlight_hex)
    latin = theme_ref + "-lt" if theme_ref else font_name
    ea = theme_ref + "-ea" if theme_ref else font_name
    f.name = latin
    ea_el = rPr.makeelement(qn("a:ea"), {"typeface": ea})
    rPr.append(ea_el)
    cs_el = rPr.makeelement(qn("a:cs"), {"typeface": theme_ref + "-cs" if theme_ref else font_name})
    rPr.append(cs_el)


# ───────────── main entry ─────────────
def build_slide(slide, src: Optional[Image.Image], background: Image.Image, pieces: List[Piece], blocks: List[Block],
                cfg, plan: FontPlan, rect_emu: Optional[Tuple[int, int, int, int]], slide_wh: Tuple[int, int],
                pdf_font_names: bool = False, orig_shape_id: Optional[int] = None):
    """Populate `slide` with the editable layers of one flattened slide.
    `orig_shape_id`: an existing picture shape (PPTX source) that is kept, hidden and lossless instead of
    being re-added as JPEG; `src` is only read when there is no such shape."""
    sw, sh = slide_wh
    if rect_emu is None:
        rect_emu = (0, 0, sw, sh)
    rl, rt, rw, rh = rect_emu
    sx, sy = rw / background.width, rh / background.height

    def X(px):  # noqa: E743
        return Emu(int(rl + px * sx))

    def Y(px):  # noqa: E743
        return Emu(int(rt + py_to_y(px)))

    def py_to_y(px):
        return px * sy

    def W(px):  # noqa: E743
        return Emu(max(1, int(px * sx)))

    def H(px):  # noqa: E743
        return Emu(max(1, int(px * sy)))

    labels = LABELS.get(cfg.ui_lang, LABELS["en"])
    existing = next((sh for sh in slide.shapes if orig_shape_id is not None and sh.shape_id == orig_shape_id), None)
    if cfg.keep_hidden_original:
        if existing is not None:
            existing.name = labels["orig"]
            _alt(existing, labels["orig"])
            _hide(existing)
        elif src is not None:
            orig = slide.shapes.add_picture(to_jpg(src), Emu(rl), Emu(rt), Emu(rw), Emu(rh))
            orig.name = labels["orig"]
            _alt(orig, labels["orig"])
            _hide(orig)
    elif existing is not None:
        existing._element.getparent().remove(existing._element)

    bg = slide.shapes.add_picture(to_jpg(upscale(background, cfg.upscale)), Emu(rl), Emu(rt), Emu(rw), Emu(rh))
    bg.name = labels["bg"]
    _alt(bg, labels["bg"])

    # pictures (back → front)
    if pieces:
        container = slide.shapes.add_group_shape() if cfg.group_shapes and len(pieces) > 1 else slide.shapes
        if container is not slide.shapes:
            container.name = labels["pics"]
        for n, pc in enumerate(pieces, 1):
            img = Image.fromarray(pc.rgba) if pc.rgba is not None else Image.open(pc.png_path)
            pic = (container.shapes if container is not slide.shapes else container).add_picture(
                to_png(upscale(img, cfg.upscale)), X(pc.x), Y(pc.y), W(pc.w), H(pc.h))
            pic.name = f"{labels['pic']} {n} · {pc.name}".strip()
            _alt(pic, pc.name)

    # text
    title_idx = None
    if cfg.title_as_placeholder:
        titles = [i for i, b in enumerate(blocks) if b.role == "title"]
        if len(titles) == 1:
            title_idx = titles[0]
    group_blocks = [i for i in range(len(blocks)) if i != title_idx]
    text_container = slide.shapes
    if cfg.group_shapes and len(group_blocks) > 1:
        text_container = slide.shapes.add_group_shape()
        text_container.name = labels["texts"]
    for n, (i, b) in enumerate(enumerate(blocks), 1):
        owner = slide.shapes if i == title_idx else (text_container.shapes if text_container is not slide.shapes else slide.shapes)
        tb = _add_text_block(owner, b, cfg, plan, X, Y, W, H, sx, sy, pdf_font_names)
        first_words = " ".join(b.text.split())[:24]
        tb.name = f"{labels['text']} {n} · {first_words}".strip(" ·")
        if i == title_idx:
            _make_title_placeholder(tb)


def _add_text_block(owner, b: Block, cfg, plan: FontPlan, X, Y, W, H, sx, sy, pdf_font_names: bool):
    x0, y0, x1, y1 = b.box
    pad = b.font_px * 0.2
    wrap = cfg.body_word_wrap and b.role == "body" and "\n" not in b.text.strip() and b.align == "left"
    slack = 0.10 if wrap else 0.0
    w_px = (x1 - x0) * (1 + slack) + 2 * pad
    tb = owner.add_textbox(X(x0 - pad), Y(y0 - pad), W(w_px), H(y1 - y0 + 2 * pad))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    size_pt = b.font_px * sy / 12700
    is_title = b.role == "title"
    fname, fbold = plan.for_block(b.bold, is_title, cfg.title_use_bold_font)
    theme_ref = None
    if cfg.use_theme_fonts:
        theme_ref = "+mj" if (is_title or (b.bold and not plan.fake_bold)) else "+mn"
    align = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(b.align, PP_ALIGN.LEFT)
    hl_hex = (b.highlight_color or "FFF59D") if b.highlights else None
    for k, para in enumerate(b.text.split("\n")):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        if b.pitch_px > 0:
            p.line_spacing = Pt(round(b.pitch_px * sy / 12700, 1))
        if b.bullet:
            _add_bullet(p, size_pt)
        for seg, hot in _segments(para, b.highlights):
            run = p.add_run()
            run.text = seg
            _style_run(run, size_pt, fbold or hot, b.color, fname, theme_ref, cfg.lang_tag,
                       hl_hex if hot else None)
    return tb


def add_failed_slide(slide, src: Image.Image, rect_emu, slide_wh, cfg):
    sw, sh = slide_wh
    rl, rt, rw, rh = rect_emu or (0, 0, sw, sh)
    pic = slide.shapes.add_picture(to_jpg(src), Emu(rl), Emu(rt), Emu(rw), Emu(rh))
    pic.name = "original"
    return pic
