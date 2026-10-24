"""Load PPTX / PDF / images into a uniform list of slides.

Slides are *lazy*: loading a deck reads only metadata (and the PDF text layer); the pixels of a slide are
rendered on first access through `SlideInput.image` and dropped again with `release()`, so a 100-page PDF
does not sit in memory all at once.

Detects which slides are really "flattened" (one picture covering the slide and no real text) so
already-editable slides pass through untouched, keeps notes and the original geometry, and reads the PDF
text layer when there is one (no OCR, no API cost).
"""
from __future__ import annotations

import io
import os
from dataclasses import dataclass, field
from typing import Callable, List, Optional, Tuple

from PIL import Image

from .types import Line

PICTURE = 13  # MSO_SHAPE_TYPE.PICTURE
EMU_PER_INCH = 914400
DEFAULT_WIDTH_EMU = 12192000  # 16:9, 13.333 in


@dataclass
class SlideInput:
    index: int                                   # 1-based position in the deck
    kind: str                                    # flat | passthrough | text_layer
    notes: str = ""
    rect_emu: Optional[Tuple[int, int, int, int]] = None   # where the picture sits (left, top, width, height)
    lines: Optional[List[Line]] = None           # PDF text layer lines (pixels of `image`)
    pic_shape_id: Optional[int] = None           # PPTX: shape id of the flattened picture
    reason: str = ""
    loader: Optional[Callable[[], dict]] = field(default=None, repr=False, compare=False)
    _data: Optional[dict] = field(default=None, repr=False, compare=False)

    @classmethod
    def from_image(cls, index: int, kind: str, image: Image.Image, **kw) -> "SlideInput":
        return cls(index, kind, loader=lambda: {"image": image}, **kw)

    def _load(self) -> dict:
        if self._data is None:
            self._data = self.loader() if self.loader else {}
        return self._data

    @property
    def image(self) -> Optional[Image.Image]:
        return self._load().get("image")

    @property
    def clean_image(self) -> Optional[Image.Image]:
        """PDF text-layer pages: the page rendered without text (and without extracted images)."""
        return self._load().get("clean_image")

    @property
    def embedded_images(self) -> List[dict]:
        return self._load().get("embedded_images", [])

    def release(self) -> None:
        self._data = None


@dataclass
class Deck:
    path: str
    kind: str                                    # pptx | pdf | image
    slides: List[SlideInput]
    width_emu: int
    height_emu: int
    source_prs: object = None                    # python-pptx Presentation when kind == pptx

    @property
    def flat_count(self) -> int:
        return sum(1 for s in self.slides if s.kind != "passthrough")


# ───────────────────────── PPTX ─────────────────────────
def _slide_has_real_text(slide) -> bool:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return True
        if getattr(sh, "has_table", False) and sh.has_table:
            return True
    return False


def _pptx_loader(blob: bytes, crop: Tuple[float, float, float, float]) -> Callable[[], dict]:
    def load():
        img = Image.open(io.BytesIO(blob)).convert("RGB")
        cl, cr, ct, cb = crop
        if any(abs(c) > 1e-4 for c in crop):
            w, h = img.size
            img = img.crop((int(w * cl), int(h * ct), int(w * (1 - cr)), int(h * (1 - cb))))
        return {"image": img}
    return load


def load_pptx(path: str, flat_cover: float = 0.85) -> Deck:
    from pptx import Presentation
    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    slides: List[SlideInput] = []
    for idx, slide in enumerate(prs.slides, 1):
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        pics = [s for s in slide.shapes if s.shape_type == PICTURE]
        if not pics:
            slides.append(SlideInput(idx, "passthrough", notes=notes, reason="no picture"))
            continue
        big = max(pics, key=lambda s: (s.width or 0) * (s.height or 0))
        cover = ((big.width or 0) * (big.height or 0)) / float(sw * sh)
        if cover < flat_cover or _slide_has_real_text(slide):
            slides.append(SlideInput(idx, "passthrough", notes=notes,
                                     reason=f"picture covers {cover:.0%}" if cover < flat_cover else "has real text"))
            continue
        try:
            blob = big.image.blob
            Image.open(io.BytesIO(blob)).verify()       # readable raster? (emf/wmf are not)
        except Exception as e:
            slides.append(SlideInput(idx, "passthrough", notes=notes, reason=f"unreadable picture: {e}"))
            continue
        crop = (big.crop_left or 0, big.crop_right or 0, big.crop_top or 0, big.crop_bottom or 0)
        slides.append(SlideInput(idx, "flat", notes=notes, rect_emu=(big.left or 0, big.top or 0, big.width, big.height),
                                 pic_shape_id=big.shape_id, loader=_pptx_loader(blob, crop)))
    return Deck(path, "pptx", slides, sw, sh, source_prs=prs)


# ───────────────────────── PDF ─────────────────────────
def _pdf_lines(page, scale: float) -> List[Line]:
    """Lines from the text layer, in rendered-pixel coordinates."""
    lines: List[Line] = []
    d = page.get_text("dict")
    bid = 0
    for blk in d.get("blocks", []):
        if blk.get("type", 0) != 0:
            continue
        bid += 1
        for ln in blk.get("lines", []):
            spans = [s for s in ln.get("spans", []) if s.get("text", "").strip()]
            if not spans:
                continue
            text = "".join(s["text"] for s in ln["spans"]).strip()
            x0 = min(s["bbox"][0] for s in spans); y0 = min(s["bbox"][1] for s in spans)
            x1 = max(s["bbox"][2] for s in spans); y1 = max(s["bbox"][3] for s in spans)
            main = max(spans, key=lambda s: len(s["text"]))
            col = main.get("color", 0)
            color = ((col >> 16) & 255, (col >> 8) & 255, col & 255)
            bold = bool(main.get("flags", 0) & 16) or "bold" in main.get("font", "").lower()
            lines.append(Line(box=(int(x0 * scale), int(y0 * scale), int(x1 * scale) + 1, int(y1 * scale) + 1),
                              text=text, conf=1.0, font_px=float(main["size"]) * scale, bold=bold,
                              color=color, font_name=main.get("font"), block_id=bid))
    return lines


def _render(page, dpi: int) -> Image.Image:
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def _pixmap_to_pil(doc, info) -> Optional[Image.Image]:
    """Embedded image → PIL, honouring its soft mask (transparency) and colour space."""
    import pymupdf
    xref, smask = info[0], info[1]
    pm = pymupdf.Pixmap(doc, xref)
    if pm.n - pm.alpha != 3:                       # gray / CMYK / indexed → RGB
        pm = pymupdf.Pixmap(pymupdf.csRGB, pm)
    if smask:
        try:
            mask = pymupdf.Pixmap(doc, smask)
            if mask.n - mask.alpha != 1:
                mask = pymupdf.Pixmap(pymupdf.csGRAY, mask)
            if (mask.width, mask.height) != (pm.width, pm.height):
                mask = None
            if mask is not None:
                pm = pymupdf.Pixmap(pm, mask)
        except Exception:
            pass
    mode = "RGBA" if pm.alpha else "RGB"
    return Image.frombytes(mode, (pm.width, pm.height), pm.samples)


def _pdf_loader(path: str, page_no: int, dpi: int, text_layer: bool, extract_images: bool) -> Callable[[], dict]:
    def load():
        import pymupdf
        scale = dpi / 72.0
        out = {}
        doc = pymupdf.open(path)
        try:
            page = doc[page_no - 1]
            out["image"] = _render(page, dpi)
            if not text_layer:
                return out
            embedded: List[dict] = []
            if extract_images:
                for info in page.get_images(full=True):
                    xref = info[0]
                    try:
                        rects = page.get_image_rects(xref)
                    except Exception:
                        rects = []
                    keep = []
                    for r in rects:
                        area = (r.width * r.height) / (page.rect.width * page.rect.height)
                        if 0.004 <= area <= 0.9:
                            keep.append(r)
                    if not keep:
                        continue                   # full-page backgrounds and specks stay in the background
                    try:
                        pil = _pixmap_to_pil(doc, info)
                    except Exception:
                        continue
                    for r in keep:
                        box = (int(r.x0 * scale), int(r.y0 * scale), int(r.x1 * scale), int(r.y1 * scale))
                        embedded.append({"box": box, "image": pil, "name": f"image {len(embedded) + 1}"})
                    try:
                        page.delete_image(xref)    # removes just this image; other images (backgrounds) survive
                    except Exception:
                        embedded = [e for e in embedded if e["image"] is not pil]
            for blk in page.get_text("dict").get("blocks", []):
                if blk.get("type", 0) == 0:
                    page.add_redact_annot(pymupdf.Rect(blk["bbox"]))
            page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
            out["clean_image"] = _render(page, dpi)
            out["embedded_images"] = embedded
        except Exception:
            out.pop("clean_image", None)
            out["embedded_images"] = []
        finally:
            doc.close()
        return out
    return load


def load_pdf(path: str, dpi: int = 150, extract_images: bool = True, min_text_chars: int = 8) -> Deck:
    import pymupdf
    doc = pymupdf.open(path)
    try:
        slides: List[SlideInput] = []
        scale = dpi / 72.0
        first = doc[0].rect if doc.page_count else None
        width_emu = DEFAULT_WIDTH_EMU
        height_emu = int(DEFAULT_WIDTH_EMU * first.height / first.width) if first else int(DEFAULT_WIDTH_EMU * 9 / 16)
        for idx, page in enumerate(doc, 1):
            lines = _pdf_lines(page, scale)
            n_chars = sum(len(l.text) for l in lines)
            if n_chars < min_text_chars:
                slides.append(SlideInput(idx, "flat", loader=_pdf_loader(path, idx, dpi, False, extract_images)))
            else:
                slides.append(SlideInput(idx, "text_layer", lines=lines,
                                         loader=_pdf_loader(path, idx, dpi, True, extract_images)))
    finally:
        doc.close()
    return Deck(path, "pdf", slides, width_emu, height_emu)


# ───────────────────────── images ─────────────────────────
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")


def load_images(paths: List[str]) -> Deck:
    slides = []
    w_emu = h_emu = None
    for idx, p in enumerate(paths, 1):
        with Image.open(p) as im:
            w, h = im.size
        if w_emu is None:
            w_emu, h_emu = DEFAULT_WIDTH_EMU, int(DEFAULT_WIDTH_EMU * h / w)
        slides.append(SlideInput(idx, "flat", loader=(lambda q=p: {"image": Image.open(q).convert("RGB")})))
    return Deck(paths[0], "image", slides, w_emu, h_emu)


def load_deck(path: str, dpi: int = 150) -> Deck:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pptx":
        return load_pptx(path)
    if ext == ".pdf":
        return load_pdf(path, dpi=dpi)
    if ext in IMAGE_EXT:
        return load_images([path])
    raise ValueError(f"Unsupported input: {path}")
