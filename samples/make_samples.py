"""Generate small synthetic inputs to try the tool without any real deck:
  samples/out/flat.png      one flattened slide (title, bullets, two shapes)
  samples/out/mixed.pptx    slide 1 flattened, slides 2-3 already editable
  samples/out/text.pdf      page 1 with a text layer + embedded image, page 2 pure raster
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

from PIL import Image, ImageDraw, ImageFont  # noqa: E402

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def draw_slide(w=1280, h=720):
    img = Image.new("RGB", (w, h), (250, 247, 240))
    d = ImageDraw.Draw(img)
    d.ellipse((900, 300, 1180, 600), fill=(60, 120, 220))
    d.rectangle((880, 120, 1200, 200), fill=(250, 200, 80))
    for text, x, y, size, fp, col in [("Polar Frontiers Report", 80, 70, 56, FONT_B, (30, 30, 60)),
                                      ("• Ice is melting faster than expected", 100, 260, 30, FONT, (40, 40, 40)),
                                      ("• Shipping routes are opening", 100, 320, 30, FONT, (40, 40, 40))]:
        d.text((x, y), text, fill=col, font=ImageFont.truetype(fp, size))
    return img


def main(out=os.path.join(HERE, "out")):
    os.makedirs(out, exist_ok=True)
    img = draw_slide()
    img.save(os.path.join(out, "flat.png"))

    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    s1.shapes.add_picture(os.path.join(out, "flat.png"), 0, 0, prs.slide_width, prs.slide_height)
    s1.notes_slide.notes_text_frame.text = "speaker notes 1"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Real editable title"
    s2.shapes.add_picture(os.path.join(out, "flat.png"), Inches(1), Inches(2), Inches(4), Inches(2))
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    s3.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "only text"
    prs.save(os.path.join(out, "mixed.pptx"))

    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=960, height=540)
    page.draw_rect(pymupdf.Rect(0, 0, 960, 540), color=None, fill=(0.97, 0.96, 0.93))
    page.insert_text((50, 80), "Deck Title From PDF", fontsize=36, fontname="hebo", color=(0.1, 0.2, 0.3))
    page.insert_text((50, 160), "• First body line in the pdf", fontsize=18, fontname="helv")
    page.insert_text((50, 190), "• Second body line in the pdf", fontsize=18, fontname="helv")
    draw_slide(400, 300).save(os.path.join(out, "emb.png"))
    page.insert_image(pymupdf.Rect(600, 250, 900, 475), filename=os.path.join(out, "emb.png"))
    page2 = doc.new_page(width=960, height=540)
    page2.insert_image(page2.rect, filename=os.path.join(out, "flat.png"))
    doc.save(os.path.join(out, "text.pdf"))
    print("samples written to", out)


if __name__ == "__main__":
    main()
